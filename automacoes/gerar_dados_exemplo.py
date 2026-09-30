# -*- coding: utf-8 -*-
"""
Gerador de planilhas de exemplo 100% fictícias para a demonstração no FluxoWMS.

Gera, com a mesma estrutura (e a mesma "bagunça") das planilhas reais:
  1) robo_cria_pedidos/dados_exemplo/RELATORIO_ENTREGAS_<data>.xlsx
     Relatório por paciente: cabeçalhos repetidos, linhas em branco, códigos e
     quantidades em colunas que mudam de lugar, números como texto, pacientes com
     "Estabelecimento" (que o robô ignora) e uma seção de CANCELADOS no fim.
  2) robo_reabastecimento/dados_exemplo/PLANILHA_DESCONTO_<data>.xlsx
     Aba "SAÍDA PROGRAMADA" com as colunas TIPO, CÓD e "PEDIDO " (com o espaço no fim,
     igual à original), centenas de linhas vazias, zeros e linhas de subtotal.

Uso:
  python gerar_dados_exemplo.py                  -> 60 pacientes
  python gerar_dados_exemplo.py --pacientes 200  -> até 200 pedidos numa execução
  python gerar_dados_exemplo.py --semente 7      -> sempre os mesmos dados

Os códigos de produto vêm de ../simulador/dados_demo.js (o mesmo catálogo do simulador).
"""
import argparse
import os
import random
import re
import sys
from datetime import datetime, timedelta

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    print("Falta o pacote openpyxl. Rode: pip install -r requirements.txt")
    sys.exit(1)

PASTA = os.path.dirname(os.path.abspath(__file__))
ARQ_CATALOGO = os.path.join(PASTA, "..", "simulador", "dados_demo.js")

PRIMEIROS_NOMES = ["ANA", "MARIA", "JOSÉ", "JOÃO", "FRANCISCA", "ANTÔNIO", "FRANCISCO", "LUCAS", "JULIANA", "PEDRO",
                   "MARCOS", "PATRÍCIA", "RAIMUNDA", "CARLOS", "SEBASTIÃO", "LETÍCIA", "BEATRIZ", "RAFAEL", "CAMILA",
                   "GABRIEL", "LARISSA", "MATEUS", "HELENA", "VICENTE", "LÚCIA", "TEREZA", "DAVI", "ISABEL", "OTÁVIO",
                   "CECÍLIA", "BENEDITO", "IRACEMA", "NATÁLIA", "SAMUEL", "VITÓRIA", "HEITOR", "ALICE", "BERNARDO"]
SOBRENOMES = ["SILVA", "SANTOS", "OLIVEIRA", "SOUSA", "LIMA", "PEREIRA", "FERREIRA", "ALVES", "RODRIGUES", "COSTA",
              "GOMES", "MARTINS", "ARAÚJO", "CARVALHO", "RIBEIRO", "BARBOSA", "CAVALCANTE", "MOREIRA", "NUNES",
              "MENDES", "FREITAS", "CARDOSO", "ROCHA", "DIAS", "TEIXEIRA", "MONTEIRO", "MOURA", "CASTRO", "BEZERRA"]
AGENDAMENTOS = ["Agendado para: {d} - Retirada no balcão", "AGENDADO p/ {d} - Entrega domiciliar",
                "Agendado: {d} (retirada)", "Agendado para {d} - Entrega na unidade"]
OBS_SOLTAS = ["Obs.: entregar em mãos ao responsável", "Obs.: conferir validade na entrega",
              "* item com troca de apresentação", "Obs.: paciente com retorno em 30 dias"]
TIPOS_REAB = ["DIETAS", "MEDICAMENTOS", "MATERIAL MÉDICO", "FRALDAS E HIGIENE", "CURATIVOS"]


def carregar_catalogo():
    if not os.path.exists(ARQ_CATALOGO):
        print(f"Catálogo não encontrado em {ARQ_CATALOGO}")
        sys.exit(1)
    texto = open(ARQ_CATALOGO, encoding="utf-8").read()
    itens = re.findall(r'\["([0-9A-Z]+)",\s*"([^"]+)",\s*"([^"]+)"\]', texto)
    if not itens:
        print("Não consegui ler os produtos do dados_demo.js")
        sys.exit(1)
    return [{"codigo": c, "descricao": d, "categoria": cat} for c, d, cat in itens]


def nome_ficticio():
    return f"{random.choice(PRIMEIROS_NOMES)} {random.choice(SOBRENOMES)} {random.choice(SOBRENOMES)}"


def registro_ficticio():
    """Registro de paciente no formato 1234567_12345678 (propositalmente diferente de qualquer numeração real)."""
    return "".join(random.choice("0123456789") for _ in range(7)) + "_" + "".join(random.choice("0123456789") for _ in range(8))


def qtd_baguncada(q):
    """A mesma quantidade pode vir como inteiro, decimal ou texto (como nas planilhas reais)."""
    r = random.random()
    if r < 0.2:
        return float(q)
    if r < 0.35:
        return str(q)
    return q


# ----------------------------------------------------------------------------------------------
# 1) Relatório por paciente (robô que cria pedidos)
# ----------------------------------------------------------------------------------------------
def gerar_relatorio(catalogo, pacientes, data_ref, destino):
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatorio"
    negrito = Font(bold=True)
    titulo = Font(bold=True, size=13, color="0B3D5C")
    cinza = PatternFill("solid", fgColor="E7ECF2")
    azul = PatternFill("solid", fgColor="DCEBFA")
    fino = Side(style="thin", color="B7C3D0")
    borda = Border(top=fino, bottom=fino, left=fino, right=fino)
    for col, larg in zip("ABCDEFGHIJ", [3, 3, 62, 11, 11, 11, 11, 16, 11, 11]):
        ws.column_dimensions[col].width = larg

    linha = [1]
    paginas = max(1, pacientes // 12 + 1)

    def escrever(valores, fonte=None, preenchimento=None, com_borda=False):
        r = linha[0]
        for col, v in valores.items():
            c = ws.cell(row=r, column=col, value=v)
            if fonte:
                c.font = fonte
            if preenchimento:
                c.fill = preenchimento
            if com_borda:
                c.border = borda
        linha[0] += 1
        return r

    def cabecalho_pagina(n):
        r = escrever({3: "RELATÓRIO DE ENTREGAS PROGRAMADAS", 8: f"Página {n} de {paginas}"}, titulo)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        escrever({3: f"Período: {data_ref:%d/%m/%Y} a {data_ref + timedelta(days=7):%d/%m/%Y}",
                  8: f"Emitido em {datetime.now():%d/%m/%Y %H:%M}"})
        linha[0] += 1

    def bloco_paciente(com_estabelecimento=False):
        variante = random.random()
        rot_reg = "REGISTRO DE PACIENTE:" if variante < 0.3 else "Registro de paciente:"
        escrever({3: f"{rot_reg} {registro_ficticio()}", 7: f"{data_ref:%d/%m}"}, negrito, cinza)
        if com_estabelecimento:
            escrever({3: f"Estabelecimento: UNIDADE PARCEIRA {random.randint(1, 9):02d}"}, None, cinza)
        rot_pac = random.choice(["Paciente: ", "Paciente:", "PACIENTE: "])
        escrever({3: rot_pac + nome_ficticio()}, negrito, cinza)
        dia = data_ref + timedelta(days=random.randint(0, 5))
        escrever({3: random.choice(AGENDAMENTOS).format(d=f"{dia:%d/%m/%Y}")})
        linha[0] += random.choice([0, 0, 1, 2])

        # a coluna do código e a distância até a quantidade mudam de paciente para paciente
        col_cod = random.choice([4, 4, 5, 6])
        col_qtd = col_cod + random.choice([1, 1, 2, 3])
        escrever({3: random.choice(["Produto", "PRODUTO / DESCRIÇÃO", "Produto (descrição)"]), col_cod: "Código", col_qtd: "Qtd."},
                 negrito, azul, True)
        escolhidos = random.sample(catalogo, random.randint(1, 12))
        for i, p in enumerate(escolhidos):
            q = random.choice([1, 1, 2, 2, 3, 4, 5, 6, 8, 10, 12, 15, 20, 30, 60])
            codigo = p["codigo"] if not p["codigo"].isdigit() or random.random() < 0.7 else int(p["codigo"])
            escrever({3: p["descricao"].upper(), col_cod: codigo, col_qtd: qtd_baguncada(q)}, None, None, True)
            if random.random() < 0.06:
                escrever({3: random.choice(OBS_SOLTAS)})
            if i == len(escolhidos) // 2 and random.random() < 0.08:
                linha[0] += 1                     # linha em branco no meio dos itens
        linha[0] += random.choice([1, 1, 2, 3])

    cabecalho_pagina(1)
    pagina = 1
    for i in range(pacientes):
        if i and i % 12 == 0:
            pagina += 1
            cabecalho_pagina(pagina)
        bloco_paciente(com_estabelecimento=random.random() < 0.08)

    linha[0] += 1
    escrever({3: "PEDIDOS CANCELADOS"}, titulo)
    for _ in range(2):
        bloco_paciente()

    nome = f"RELATORIO_ENTREGAS_{data_ref:%d_%m_%Y}.xlsx"
    os.makedirs(destino, exist_ok=True)
    caminho = os.path.join(destino, nome)
    wb.save(caminho)
    return caminho


# ----------------------------------------------------------------------------------------------
# 2) Planilha de desconto / reabastecimento (aba SAÍDA PROGRAMADA)
# ----------------------------------------------------------------------------------------------
def gerar_planilha_desconto(catalogo, data_ref, destino):
    wb = Workbook()
    ws = wb.active
    ws.title = "SAÍDA PROGRAMADA"
    colunas = ["DESCRIÇÃO", "CÓD", "QTD", "TIPO", "ESTOQ. WMS UNIDADE", "SITUAÇÃO", "SAÍDAS DO DIA",
               "QTD ATD DIA", "ESTOQ. DISP. CD", "PEDIDO ", "PÓS PEDIDO", "FATOR EMBALAG"]
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D7A66")
    for col, larg in zip("ABCDEFGHIJKL", [52, 10, 8, 20, 18, 11, 13, 12, 15, 10, 11, 14]):
        ws.column_dimensions[col].width = larg

    produtos = [p for p in catalogo if p["categoria"] in TIPOS_REAB]
    estoques = []
    for tipo in TIPOS_REAB:
        do_tipo = [p for p in produtos if p["categoria"] == tipo]
        soma = 0
        for p in do_tipo:
            qtd = random.randint(5, 120)
            estoque = random.randint(0, 70)
            atend = random.randint(0, 40)
            disp_cd = random.randint(50, 900)
            fator = random.choice([1, 1, 6, 10, 12, 24, 50])
            precisa = max(0, qtd + atend - estoque)
            pedido = (precisa // fator + (1 if precisa % fator else 0)) * fator if precisa else 0
            if random.random() < 0.10:
                pedido = 0                      # itens sem necessidade de pedido
            if pedido == 0 and random.random() < 0.2:
                pedido = None                   # célula vazia
            if pedido is None:
                valor_pedido = None
            else:
                valor_pedido = float(pedido)
            codigo = int(p["codigo"]) if p["codigo"].isdigit() else p["codigo"]
            ws.append([p["descricao"].upper(), codigo, qtd, tipo, estoque, estoque - qtd, atend, atend,
                       disp_cd, valor_pedido, (estoque + (pedido or 0)) - qtd, f"CX {fator}" if fator > 1 else "UN"])
            estoques.append((codigo, p["descricao"], estoque, disp_cd, fator))
            soma += pedido or 0
        # linha de subtotal (sem código e sem TIPO: o robô precisa ignorar)
        ws.append([f"SUBTOTAL {tipo}", None, None, None, None, None, None, None, None, float(soma), None, None])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, italic=True)
        ws.append([])                            # linha em branco entre grupos

    # centenas de linhas "vazias" com fórmulas arrastadas que dão zero (como na original)
    for _ in range(1100):
        ws.append([None, None, None, None, None, None, None, None, None, 0.0, 0.0, None])

    aba = wb.create_sheet("WMS_REL.SINTETICO.PROD")
    aba.append(["", "Código", "", "", "", "", "Estoque"])
    for cod, desc, est, _, _ in estoques:
        aba.append(["", cod, desc, "", "", "", est])
    aba = wb.create_sheet("BASE SAÍDAS DO DIA")
    aba.append(["CÓD", "DESCRIÇÃO", "QTD ATENDIDA", "DATA"])
    for cod, desc, _, _, _ in random.sample(estoques, min(30, len(estoques))):
        aba.append([cod, desc, random.randint(1, 30), f"{data_ref:%d/%m/%Y}"])
    aba = wb.create_sheet("ESTOQUE CD CENTRAL")
    aba.append(["", "Código", "", "", "", "", "Disponível"])
    for cod, desc, _, disp, _ in estoques:
        aba.append(["", cod, desc, "", "", "", disp])
    aba = wb.create_sheet("FATOR EMB")
    aba.append(["CÓD", "DESCRIÇÃO", "EMBALAGEM", "FATOR"])
    for cod, desc, _, _, fator in estoques:
        aba.append([cod, desc, "CX" if fator > 1 else "UN", fator])

    nome = f"PLANILHA_DESCONTO_{data_ref:%d_%m_%Y}.xlsx"
    os.makedirs(destino, exist_ok=True)
    caminho = os.path.join(destino, nome)
    wb.save(caminho)
    return caminho


def main():
    ap = argparse.ArgumentParser(description="Gera planilhas fictícias para a demonstração no FluxoWMS.")
    ap.add_argument("--pacientes", type=int, default=60, help="quantidade de pacientes no relatório (padrão 60)")
    ap.add_argument("--semente", type=int, default=None, help="semente aleatória (para repetir os mesmos dados)")
    ap.add_argument("--data", default=None, help="data de referência DD/MM/AAAA (padrão: hoje)")
    args = ap.parse_args()

    if args.semente is not None:
        random.seed(args.semente)
    data_ref = datetime.strptime(args.data, "%d/%m/%Y") if args.data else datetime.now()
    catalogo = carregar_catalogo()

    rel = gerar_relatorio(catalogo, max(1, args.pacientes), data_ref, os.path.join(PASTA, "robo_cria_pedidos", "dados_exemplo"))
    des = gerar_planilha_desconto(catalogo, data_ref, os.path.join(PASTA, "robo_reabastecimento", "dados_exemplo"))
    print("Planilhas fictícias geradas:")
    print("  " + rel)
    print("  " + des)


if __name__ == "__main__":
    main()
